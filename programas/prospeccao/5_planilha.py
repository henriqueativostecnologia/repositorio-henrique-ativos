# Passo 5: planilha de visitas (abas Lojas, Rotas no Maps, Resumo do teste e Como ler).
# Uso: python3 5_planilha.py dados/roteiro.json dados/Roteiro_CIDADE_DATA.xlsx
# Tem dados reais: o arquivo gerado nunca entra no repositório.
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
R = json.load(open(sys.argv[1])); L = R['L']; OUT = sys.argv[2]; CF = R['config']
RES = ['Não falei com quem decide', 'Sem interesse', 'Talvez', 'Interessado', 'Fechou']
H = ['Parada', 'Bloco', 'Tipo de parada', 'Grupo: com site / sem site', 'Segmento', 'Nome', 'Bairro', 'Endereço', 'Telefone', 'Alertas',
     'Nota', 'Avaliações', 'Ficha sem dono', 'Situação do site', 'Site atual', 'Rede/franquia', 'Gourmet (só informação)', 'Melhor posição (só informação)',
     'Faixa de nota', 'Faixa de avaliações', 'Km desde a parada anterior (linha reta)', 'Tipo no Google', 'Achada pelos termos', 'Link do Maps',
     'Resultado', 'Data da visita', 'Falou com quem', 'Objeção', 'Próximo passo']
C = {h: i + 1 for i, h in enumerate(H)}; col = lambda h: CL(C[h])
wb = Workbook(); ws = wb.active; ws.title = 'Lojas'; ws.append(H)
CAB = PatternFill('solid', fgColor='1D2B27'); BL = PatternFill('solid', fgColor='E9F0EC'); OP = PatternFill('solid', fgColor='EFEFEF')
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = CAB; c.alignment = Alignment(wrap_text=True, vertical='center')
ws.row_dimensions[1].height = 45
for l in L:
    op = not l['plan']
    ws.append([l['parada'] + (' (opcional)' if op else ''), l['bloco_nome'], 'Opcional' if op else 'Planejada', l['grupo2'], l['seg'], l['nome'], l['bairro'],
               l['end'], l['tel'], '; '.join(l['alertas']), l['nota'], l['aval'], 'Sim' if l['sem_dono'] else 'Não', l['sit'], l['site'],
               'Sim' if l['rede'] else '', l['gourmet'], l['pos'], l['fnota'], l['faval'], round(l['dist'], 2), l['tipo'], ', '.join(l['qs']), l['link'],
               None, None, None, None, None])
    n = ws.max_row
    for c in ws[n]:
        c.alignment = Alignment(vertical='top', wrap_text=c.column_letter in (col('Nome'), col('Endereço'), col('Alertas'), col('Bloco')))
        if op: c.font = Font(color='6B6B6B')
        if l['bloco'] % 2 == 0 and c.column <= C['Link do Maps']: c.fill = BL
    ws.cell(n, C['Link do Maps']).hyperlink = l['link']
    ws.cell(n, C['Nota']).number_format = '0.0'; ws.cell(n, C['Avaliações']).number_format = '#,##0'; ws.cell(n, C['Km desde a parada anterior (linha reta)']).number_format = '0.00'
W = [11, 26, 10, 11, 22, 30, 18, 38, 16, 34, 6, 10, 9, 18, 26, 8, 10, 10, 12, 12, 10, 22, 22, 26, 22, 11, 16, 24, 24]
for i, w in enumerate(W, 1): ws.column_dimensions[CL(i)].width = w
ws.freeze_panes = 'G2'; ws.auto_filter.ref = ws.dimensions
last = ws.max_row
dv = DataValidation(type='list', formula1='"' + ','.join(RES) + '"', allow_blank=True, showErrorMessage=True,
                    errorTitle='Resultado', error='Escolha um item da lista.')
ws.add_data_validation(dv); dv.add(f"{col('Resultado')}2:{col('Resultado')}{last}")
# rotas
Rt = wb.create_sheet('Rotas no Maps'); Rt.append(['Bloco', 'Planejadas', 'Opcionais perto', 'Km do bloco (linha reta)', 'Rota no Google Maps (sai de onde você está e passa pelas planejadas em ordem)'])
for r in R['rotas']:
    Rt.append([r['nome'], r['n'], r['nop'], round(r['km'], 1)] + r['links'])
    for j, u in enumerate(r['links']): Rt.cell(Rt.max_row, 5 + j).hyperlink = u
for c in Rt[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = CAB; c.alignment = Alignment(wrap_text=True)
for i, w in enumerate([40, 11, 11, 11, 60, 60], 1): Rt.column_dimensions[CL(i)].width = w
# resumo do teste
S = wb.create_sheet('Resumo do teste')
S.append(['Recorte', 'Valor', 'Lojas', 'Com resultado', 'Falou com quem decide', 'Interessado + Fechou', 'Fechou', '% Interessado + Fechou (de quem falou com quem decide)'])
for c in S[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = CAB; c.alignment = Alignment(wrap_text=True, vertical='center')
S.row_dimensions[1].height = 45
rng = lambda h: f"Lojas!${col(h)}$2:${col(h)}${last}"
RE = rng('Resultado')
def bloco(titulo, h, valores):
    S.append([titulo]); S.cell(S.max_row, 1).font = Font(bold=True)
    for v in valores:
        r = S.max_row + 1; cond = f'{rng(h)},$B{r}' if h else None
        f = (lambda extra: f'=COUNTIFS({cond},{extra})' if cond else f'=COUNTIFS({extra})')
        lojas = f'=COUNTIFS({cond})' if cond else f'=COUNTA({rng("Nome")})'
        S.append(['', v, lojas, f(f'{RE},"<>"'), f'{f(f"{RE},\"<>\"")}-{f(f"{RE},\"Não falei com quem decide\"")[1:]}',
                  f'{f(f"{RE},\"Interessado\"")}+{f(f"{RE},\"Fechou\"")[1:]}', f(f'{RE},"Fechou"'), f'=IF(E{r}=0,"",F{r}/E{r})'])
        S.cell(r, 8).number_format = '0%'
bloco('Total', None, ['todas as lojas'])
bloco('Grupo', 'Grupo: com site / sem site', ['com site', 'sem site'])
bloco('Segmento', 'Segmento', sorted({l['seg'] for l in L}))
bloco('Ficha sem dono', 'Ficha sem dono', ['Sim', 'Não'])
bloco('Faixa de nota', 'Faixa de nota', ['4,5 a 5,0', '4,0 a 4,4', 'abaixo de 4,0', 'sem nota'])
bloco('Faixa de avaliações', 'Faixa de avaliações', ['200 ou mais', '50 a 199', '10 a 49', 'menos de 10'])
bloco('Tipo de parada', 'Tipo de parada', ['Planejada', 'Opcional'])
for i, w in enumerate([20, 38, 9, 12, 14, 14, 9, 22], 1): S.column_dimensions[CL(i)].width = w
S.freeze_panes = 'A2'
# como ler
T = wb.create_sheet('Como ler')
for t in [f"Prospecção {CF['titulo']}: {CF['area_descricao']}. Fonte: {CF['fonte']}.",
          'Objetivo: vender. O teste com site × sem site vai embutido: nenhuma loja foi cortada para igualar os grupos; o resumo compara a % de "Interessado + Fechou" de cada grupo.',
          'Planejada = 10 ou mais avaliações e nota 4,0 ou mais, ou ficha sem dono. Opcional = as demais (cinza). Visite todas as planejadas na ordem; as opcionais têm o número da planejada mais próxima com uma letra (ex.: 12a).',
          f'Blocos: até 10 planejadas, do mais perto para o mais longe do ponto de saída ({CF['ponto_saida']['nome']}, ponto aproximado, sem número). A aba "Rotas no Maps" abre cada bloco no Google Maps saindo de onde você está, com as 10 planejadas em ordem (as opcionais não entram na rota); a distância em linha reta é menor que a de carro.',
          'Grupo "sem site" inclui quem não tem site ou cujo "site" na ficha é só Instagram, Facebook, WhatsApp, iFood, página de links ou link encurtado (coluna "Situação do site").',
          'Alertas: "Ficha sem dono: mostrar no celular" (o Google mostra "Reivindicar esta empresa"); "Poucas avaliações: confirme que está aberta"; "Nota baixa: conversa sobre reputação"; "Rede: perguntar quem decide o marketing"; "Conferir se é do segmento"; "Conferir se é a mesma loja da parada X" (duas fichas da mesma marca).',
          'Resultado: escolha na lista (Não falei com quem decide · Sem interesse · Talvez · Interessado · Fechou). A aba "Resumo do teste" calcula sozinha; a % considera só as visitas em que você falou com quem decide.',
          'Fora da lista: ' + (', '.join(CF['fora_do_roteiro']) or 'nenhuma loja') + '. Gourmet e Melhor posição são só informação.',
          'Dados reais: este arquivo não vai para o repositório.']:
    T.append([t]); T.cell(T.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
T.column_dimensions['A'].width = 130
wb.save(OUT); print('ok', len(L), last)
