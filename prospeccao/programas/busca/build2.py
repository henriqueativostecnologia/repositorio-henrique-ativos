import sys,json,math;sys.path.insert(0,'.')
from bh import *
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
R=load(lambda p: p.get('engine')=='google_maps' and str(p.get('ll','')).endswith('14z'))
EX2=EXCL|{'Departamento de veículos automotores','Depósito','Estacionamento','Oficina mecânica'}
def seg(r):
  t=(r.get('type') or '').lower(); q=' '.join(r['_q'])
  if any(k in t for k in ('revendedora','concession','agente comercial de carros')) or (t=='loja' and 'seminovos' in q): return 'Revenda de seminovos'
  if any(k in t for k in ('martelinho','funilaria','pintura')): return 'Martelinho / funilaria'
  if any(k in t for k in ('película','envelopamento','vidros','acessórios')): return 'Insulfilm / envelopamento'
  if any(k in t for k in ('limpeza completa','detalhamento')): return 'Estética / detailing'
  if 'lava' in t: return 'Lava jato'
  return 'Estética / detailing' if not t else 'Outro'
L=[]
for r in R.values():
  b=bairro(r.get('address') or ''); g=regional(b)
  if g not in ('Pampulha','Venda Nova') or r.get('type') in EX2: continue
  s=seg(r)
  if s=='Outro': continue
  f=flags(r); rt=r.get('rating'); rv=r.get('reviews')
  f['melhor_novo']=int((f['sem_site'] or f['social']) and rt is not None and rt>=4.5)
  f['ativo']=int(rv is not None and rv>=10)
  pts=3*f['sem_site']+2*f['social']+f['poucas']+2*f['melhor_novo']-10*f['cuidado']
  mot=[]
  if f['melhor_novo']: mot.append('MELHOR ALVO: '+('sem site' if f['sem_site'] else 'só rede social')+' e nota ≥ 4,5')
  else:
    if f['sem_site']: mot.append('sem site')
    if f['social']: mot.append('"site" é só Instagram/Linktree/Facebook')
  if f['poucas']: mot.append('sem avaliações' if rv is None else f'poucas avaliações ({rv})')
  if f['cuidado']: mot.append(f'CUIDADO: nota {rt} abaixo de 4,0')
  if not f['fraca']: mot.append('presença boa: alvo difícil')
  L.append(dict(r=r,b=b,g=g,seg=s,f=f,pts=pts,mot='; '.join(mot)))
key=lambda x:(-x['f']['melhor_novo'],-x['f']['ativo'],-x['pts'],-(x['r'].get('rating') or 0),-(x['r'].get('reviews') or 0))
EST=sorted([x for x in L if x['seg']!='Revenda de seminovos' and x['f']['fraca']],key=key)
REV=sorted([x for x in L if x['seg']=='Revenda de seminovos' and x['f']['fraca']],key=key)
FORTES=[x for x in L if not x['f']['fraca']]
# roteiro: melhores alvos de estética, saindo do Jaraguá
HOME=(-19.8600,-43.9550)
def km(a,b):
  dx=(a[1]-b[1])*111.320*math.cos(math.radians(-19.85)); dy=(a[0]-b[0])*110.574; return math.hypot(dx,dy)
pos=lambda x:(x['r']['gps_coordinates']['latitude'],x['r']['gps_coordinates']['longitude'])
alvo=[x for x in EST if x['f']['melhor_novo'] and x['f']['ativo']]
route=[];cur=HOME;rest=alvo[:]
while rest:
  n=min(rest,key=lambda x:km(cur,pos(x))); route.append(n); rest.remove(n); cur=pos(n)
def tot(rt): return km(HOME,pos(rt[0]))+sum(km(pos(rt[i]),pos(rt[i+1])) for i in range(len(rt)-1))+km(pos(rt[-1]),HOME)
imp=True
while imp:
  imp=False
  for i in range(len(route)-1):
    for j in range(i+1,len(route)):
      nr=route[:i]+route[i:j+1][::-1]+route[j+1:]
      if tot(nr)<tot(route)-1e-9: route=nr; imp=True
json.dump(dict(n_all=len(L),est=len(EST),rev=len(REV),fortes=len(FORTES),route=[x['r']['title'] for x in route],tot=tot(route) if route else 0,
  cnt={g:{'melhor_ativo':sum(1 for x in EST if x['g']==g and x['f']['melhor_novo'] and x['f']['ativo']),'melhor':sum(1 for x in EST if x['g']==g and x['f']['melhor_novo']),'fracas':sum(1 for x in EST if x['g']==g)} for g in ('Pampulha','Venda Nova')}),open('etapa2_resumo.json','w'),ensure_ascii=False)
# planilha
wb=Workbook(); F=Font(name='Arial',size=10); B=Font(name='Arial',size=10,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78'); GOOD=PatternFill('solid',fgColor='E2EFDA'); BAD=PatternFill('solid',fgColor='FCE4D6')
H=['Ordem','Melhor alvo','Ativo (10+ aval.)','Nome','Segmento','Tipo no Google','Regional','Bairro','Endereço','Telefone','Site','Nota','Avaliações','Sem site','Só rede social','<30 aval.','Pontos','Motivo','Link do Maps']
def sheet(ws,data):
  ws.append(H)
  for c in ws[1]: c.font=B;c.fill=HF;c.alignment=Alignment(wrap_text=True,vertical='center')
  for i,x in enumerate(data,1):
    r=x['r'];f=x['f'];n=i+1
    ws.append([i,'SIM' if f['melhor_novo'] else '','SIM' if f['ativo'] else '',r['title'],x['seg'],r.get('type'),x['g'],x['b'],r.get('address'),r.get('phone'),r.get('website'),r.get('rating'),r.get('reviews'),f['sem_site'] or None,f['social'] or None,f['poucas'] or None,
      f'=3*N(N{n})+2*N(O{n})+1*N(P{n})+IF(AND(OR(N(N{n})=1,N(O{n})=1),N(L{n})>=4.5),2,0)-IF(AND(L{n}<>"",N(L{n})<4),10,0)',x['mot'],f"https://www.google.com/maps/place/?q=place_id:{r['place_id']}"])
    for c in ws[n]: c.font=F; c.alignment=Alignment(vertical='top',wrap_text=c.column_letter in ('D','I','R'))
    ws.cell(n,19).hyperlink=ws.cell(n,19).value
    fill=BAD if f['cuidado'] else (GOOD if f['melhor_novo'] else None)
    if fill:
      for c in ws[n]: c.fill=fill
  for i,w in enumerate([7,9,9,32,20,22,11,18,42,16,28,7,10,8,9,8,8,44,28],1): ws.column_dimensions[get_column_letter(i)].width=w
  ws.freeze_panes='E2'; ws.auto_filter.ref=ws.dimensions
ws=wb.active; ws.title='Alvos estética'; sheet(ws,EST)
sheet(wb.create_sheet('Revendas seminovos'),REV)
w3=wb.create_sheet('Roteiro Jaraguá')
w3.append(['Parada','Nome','Segmento','Bairro','Endereço','Telefone','Nota','Avaliações','Motivo','Km desde a parada anterior (linha reta)','Km acumulado (linha reta)','Link do Maps'])
for c in w3[1]: c.font=B;c.fill=HF;c.alignment=Alignment(wrap_text=True,vertical='center')
w3.append([0,'Saída: Jaraguá (ponto aproximado do bairro)','','Jaraguá','','','','','','',0,f'https://www.google.com/maps/search/?api=1&query={HOME[0]},{HOME[1]}'])
prev=HOME
for i,x in enumerate(route,1):
  r=x['r'];p=pos(x);n=i+2
  w3.append([i,r['title'],x['seg'],x['b'],r.get('address'),r.get('phone'),r.get('rating'),r.get('reviews'),x['mot'],km(prev,p),f'=K{n-1}+J{n}',f"https://www.google.com/maps/place/?q=place_id:{r['place_id']}"]); prev=p
n=w3.max_row+1
w3.append(['Volta','Retorno ao Jaraguá','','','','','','','',km(prev,HOME),f'=K{n-1}+J{n}',''])
for row in w3.iter_rows(min_row=2):
  for c in row: c.font=F; c.alignment=Alignment(vertical='top',wrap_text=c.column_letter in ('B','E','I'))
  row[9].number_format='0.00'; row[10].number_format='0.00'
  if row[11].value: row[11].hyperlink=row[11].value
for i,w in enumerate([7,32,20,18,42,16,7,10,40,12,12,28],1): w3.column_dimensions[get_column_letter(i)].width=w
# trechos com link de rota do Google Maps (até 9 paradas por trecho)
w3.append([]); w3.append(['Rotas no Google Maps (abrem a rota de carro, em trechos de até 9 paradas)']); w3.cell(w3.max_row,1).font=Font(name='Arial',size=10,bold=True)
seq=[HOME]+[pos(x) for x in route]+[HOME]
k=0
while k<len(seq)-1:
  part=seq[k:k+10]; url='https://www.google.com/maps/dir/'+'/'.join(f'{a},{b}' for a,b in part)
  w3.append([f'Trecho {k//9+1}',f'paradas {k} a {min(k+9,len(seq)-1) if k+9<len(seq)-1 else "volta"}',url]); w3.cell(w3.max_row,3).hyperlink=url; w3.cell(w3.max_row,3).font=F; k+=9
w4=wb.create_sheet('Legenda')
for t in ['Como ler esta planilha',
 'Área: regionais Pampulha e Venda Nova (lista oficial de bairros da PBH). Fonte: Google Maps via SerpApi, 15/10/2026.',
 'Buscas: grade de 9 pontos (3 km); em cada ponto, página 1 de "estética automotiva", "lava jato", "detailing polimento vitrificação", "insulfilm envelopamento", "martelinho de ouro" e "revenda de seminovos"; páginas 2 e 3 de estética e lava jato. Todas vieram cheias (20 de 20): ainda existem mais lojas além destas.',
 'O Maps mostra primeiro as lojas mais fortes; as mais fracas tendem a ficar nas páginas seguintes, que não foram buscadas.',
 'Melhor alvo (coluna B) = sem site OU "site" que é só Instagram/Linktree/Facebook, E nota 4,5 ou mais. Ativo (coluna C) = 10 avaliações ou mais.',
 'Ordem: Melhor alvo primeiro; depois Ativo; depois Pontos; depois nota maior; depois mais avaliações.',
 'Pontos (fórmula): sem site 3; só rede social 2; menos de 30 avaliações 1; nota ≥ 4,5 com site fraco +2; nota abaixo de 4,0 −10 (linha laranja, fim da lista).',
 'Nas abas de alvos entram só lojas com presença fraca (sem site, só rede social ou menos de 30 avaliações). Lojas com site próprio e 30+ avaliações ficaram de fora.',
 'Roteiro: só os Melhores alvos ativos de estética (sem revendas), na ordem que dá o menor caminho em linha reta, saindo de um ponto aproximado do Jaraguá e voltando. A distância real de carro é maior: use os links de rota.',
 'Nota vazia = o Google não mostra nota (não é zero).',
 'Dados reais: este arquivo não entra no cofre.']:
  w4.append([t]); c=w4.cell(w4.max_row,1); c.font=Font(name='Arial',size=10,bold=(w4.max_row==1)); c.alignment=Alignment(wrap_text=True)
w4.column_dimensions['A'].width=120
wb.save('Alvos_Pampulha_VendaNova_automotivo_2026-10-15.xlsx')
json.dump({'est':[(x['r']['title'],x['pts']) for x in EST],'rev':[(x['r']['title'],x['pts']) for x in REV]},open('chk2.json','w'),ensure_ascii=False)
print(open('etapa2_resumo.json').read())
