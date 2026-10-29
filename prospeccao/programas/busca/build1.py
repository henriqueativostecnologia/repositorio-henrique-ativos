import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
J=json.load(open('etapa1.json')); REG=['Barreiro','Centro-Sul','Leste','Nordeste','Noroeste','Norte','Oeste','Pampulha','Venda Nova']
F=Font(name='Arial',size=10); B=Font(name='Arial',size=10,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78')
def hdr(ws):
  for c in ws[1]: c.font=B;c.fill=HF;c.alignment=Alignment(wrap_text=True,vertical='center')
def body(ws,widths):
  for row in ws.iter_rows(min_row=2):
    for c in row:
      if c.font!=B: c.font=F
  for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
wb=Workbook(); ws=wb.active; ws.title='Ranking regionais'
ws.append(['Regional','Pontos da grade (aprox.)','Lojas na amostra','Sem site','"Site" só Instagram/Linktree/Facebook','Menos de 30 avaliações','Presença fraca (lojas)','% presença fraca','Posição pelo nº de presença fraca','Posição pelo %','Posição média (menor = melhor)'])
pts={r['g']:r['pts'] for r in J['rows']}
order=sorted(REG,key=lambda g:-next(r['fraca'] for r in J['rows'] if r['g']==g))
for i,g in enumerate(order,2):
  L="Lojas!$C:$C"
  ws.append([g,pts.get(g,0),f'=COUNTIFS({L},A{i})',f'=COUNTIFS({L},A{i},Lojas!$J:$J,1)',f'=COUNTIFS({L},A{i},Lojas!$K:$K,1)',f'=COUNTIFS({L},A{i},Lojas!$L:$L,1)',f'=COUNTIFS({L},A{i},Lojas!$M:$M,1)',
    f'=IF(C{i}=0,"",G{i}/C{i})',f'=RANK(G{i},$G$2:$G$10)',f'=RANK(H{i},$H$2:$H$10)',f'=(I{i}+J{i})/2'])
n=len(order)+2
ws.append(['A conferir','',f'=COUNTIFS(Lojas!$C:$C,A{n})',f'=COUNTIFS(Lojas!$C:$C,A{n},Lojas!$J:$J,1)',f'=COUNTIFS(Lojas!$C:$C,A{n},Lojas!$K:$K,1)',f'=COUNTIFS(Lojas!$C:$C,A{n},Lojas!$L:$L,1)',f'=COUNTIFS(Lojas!$C:$C,A{n},Lojas!$M:$M,1)',f'=IF(C{n}=0,"",G{n}/C{n})','','',''])
ws.append(['Total','=SUM(B2:B10)',f'=SUM(C2:C{n})',f'=SUM(D2:D{n})',f'=SUM(E2:E{n})',f'=SUM(F2:F{n})',f'=SUM(G2:G{n})',f'=G{n+1}/C{n+1}','','',''])
hdr(ws); body(ws,[34,12,11,9,16,12,13,11,13,11,13])
for r in range(2,n+2): ws.cell(r,8).number_format='0.00%'
for c in ws[n+1]: c.font=Font(name='Arial',size=10,bold=True)
ws.freeze_panes='B2'
notes=['LEIA ANTES DE USAR O RANKING',
'1) O número absoluto é da AMOSTRA, não o total real de lojas da regional. Foram 2 buscas ("estética automotiva" e "lava jato") em cada ponto de uma grade de 3 km (43 pontos), e o Google Maps devolve no máximo 20 lojas por busca. Todos os pontos vieram cheios (20 de 20), então há mais lojas do que as contadas.',
'2) O Maps mostra primeiro as lojas mais fortes (mais avaliações, ficha mais completa). Por isso, nas áreas cheias, o % de presença fraca tende a sair MENOR do que o real: as lojas fracas ficam para as páginas seguintes, que não foram buscadas nesta etapa.',
'3) Pontos da grade por regional: cada ponto foi atribuído à regional que mais apareceu nos resultados dele. Regionais com poucos pontos (Centro-Sul e Noroeste com 2, Leste com 0) foram menos amostradas; as lojas delas apareceram pelos pontos vizinhos.',
'4) Presença fraca = sem site, OU "site" que é só Instagram/Linktree/Facebook, OU menos de 30 avaliações (sem avaliações conta como menos de 30). Uma loja pode ter mais de um motivo; a coluna G conta cada loja uma vez só.',
'5) Regional pelo bairro do endereço do Google, conferido com a lista oficial da PBH (Relação de bairros, regionais e territórios de gestão compartilhada, mapaculturalbh.pbh.gov.br, 487 bairros). Bairro que não está na lista, ou que o endereço põe em outra regional, fica "A conferir".',
'6) Ficaram de fora: lojas fora de Belo Horizonte (Contagem, Nova Lima etc.) e categorias que não são do segmento (mecânica, estacionamento, depósito, etc.).',
'Fonte: Google Maps via SerpApi, 15/10/2026. Buscas desta etapa: 86.']
ws.append([]); 
for t in notes:
  ws.append([t]); r=ws.max_row; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=11); c=ws.cell(r,1); c.font=Font(name='Arial',size=10,bold=t.startswith('LEIA')); c.alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[r].height=15 if t.startswith('LEIA') else 42
w2=wb.create_sheet('Lojas')
w2.append(['Nome','Tipo no Google','Regional','Bairro (do endereço)','Endereço','Telefone','Site','Nota','Avaliações','Sem site','Só rede social','<30 avaliações','Presença fraca','Link do Maps'])
for x in sorted(J['L'],key=lambda x:(x['g'],-x['f']['fraca'],x['nome'])):
  r=x['r'];f=x['f']
  w2.append([x['nome'],r['type'],x['g'],x['b'] or '(sem bairro no endereço)',r['address'],r['phone'],r['website'],r['rating'],r['reviews'],f['sem_site'],f['social'],f['poucas'],f['fraca'],f"https://www.google.com/maps/place/?q=place_id:{r['place_id']}"])
hdr(w2); body(w2,[34,24,13,22,46,16,30,7,10,8,9,9,9,30]); w2.freeze_panes='B2'; w2.auto_filter.ref=w2.dimensions
for row in w2.iter_rows(min_row=2): row[13].hyperlink=row[13].value
wb.save('Ranking_regionais_BH_automotivo_2026-10-15.xlsx')
