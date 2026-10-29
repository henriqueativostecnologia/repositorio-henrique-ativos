import sys,re;sys.path.insert(0,'.')
from load import allres; from geo import dist_av
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
R=allres()
SOCIAL=('instagram.com','linktr.ee','facebook.com','fb.com','wa.me','web.app')
CHAIN=('Rede Exemplo A','Rede Exemplo B')  # exemplos: trocar pelos nomes de redes/franquias da região
def cat(t):
  t=t.lower()
  for k,v in [('padaria','Padaria'),('café','Cafeteria'),('cafeteria','Cafeteria'),('confeitaria','Cafeteria'),('lanchonete','Cafeteria'),('bar','Bar'),('gastropub','Bar'),('supermercado','Armazém/mercado'),('empório','Armazém/mercado'),('lava','Estética automotiva'),('automotiv','Estética automotiva'),('concession','Compra e venda de veículos'),('carros','Compra e venda de veículos')]:
    if k in t: return v
  return 'Restaurante'
def flags(r):
  w=(r.get('website') or '').lower(); rv=r.get('reviews'); rt=r.get('rating')
  f={'sem_site':int(not w),'social':int(bool(w) and any(s in w for s in SOCIAL)),
     'poucas':int(rv is None or rv<30),'cuidado':int(rt is not None and rt<4.0)}
  fraca=f['sem_site'] or f['social'] or f['poucas']
  f['melhor']=int(rt is not None and rt>=4.5 and bool(fraca))
  mot=[]
  if f['melhor']: mot.append('MELHOR ALVO: nota alta com presença fraca')
  if f['sem_site']: mot.append('sem site')
  if f['social']: mot.append('"site" é só '+('página de reserva' if 'web.app' in w else 'Instagram/Linktree/Facebook'))
  if f['poucas']: mot.append('sem avaliações' if rv is None else f'poucas avaliações ({rv})')
  if f['cuidado']: mot.append(f'CUIDADO: nota {rt} abaixo de 4,0')
  if not mot: mot.append('presença boa (site próprio e muitas avaliações): alvo difícil')
  if any(c.lower() in r['title'].lower() for c in CHAIN): mot.append('rede/franquia: decisão pode não ser local')
  return f,'; '.join(mot)
W={'sem_site':3,'social':2,'poucas':1,'melhor':2,'cuidado':-10}
def pts(f): return sum(W[k]*f[k] for k in W)
rows=[];extra=[]
for r in R.values():
  g=r['gps_coordinates'];d,s=dist_av(g['latitude'],g['longitude']);a=r.get('address') or ''
  if not ('Vila da Serra' in a or 'Vale do Sereno' in a) or a.startswith('Rod.'): 
    if ('Vila da Serra' in a or 'Vale do Sereno' in a) and d<=1000 and cat(r['type']) in ('Estética automotiva','Compra e venda de veículos'): extra.append((d,s,r))
    continue
  if d<=300: rows.append((d,s,r))
  elif d<=1000 and cat(r['type']) in ('Estética automotiva','Compra e venda de veículos'): extra.append((d,s,r))
def rec(d,s,r):
  f,m=flags(r)
  return dict(nome=r['title'],tipo=cat(r['type']),cat_google=r['type'],end=r.get('address'),tel=r.get('phone'),site=r.get('website'),nota=r.get('rating'),aval=r.get('reviews'),f=f,mot=m,pts=pts(f),dist=round(d),pos=round(s),
    onav='Oscar Niemeyer' in (r.get('address') or '') or 'Niemyer' in (r.get('address') or ''),link=f"https://www.google.com/maps/place/?q=place_id:{r['place_id']}")
L=[rec(*x) for x in rows]; E=[rec(*x) for x in extra]
key=lambda x:(-x['pts'],-(x['nota'] or 0),-(x['aval'] or 0))
L.sort(key=key); E.sort(key=key)
for i,x in enumerate(L,1): x['rank']=i
wb=Workbook(); F=Font(name='Arial',size=10); B=Font(name='Arial',size=10,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78')
GOOD=PatternFill('solid',fgColor='E2EFDA'); BAD=PatternFill('solid',fgColor='FCE4D6')
def sheet(ws,hdr,data,widths):
  ws.append(hdr)
  for c in ws[1]: c.font=B;c.fill=HF;c.alignment=Alignment(wrap_text=True,vertical='center')
  for row in data: ws.append(row)
  for row in ws.iter_rows(min_row=2):
    for c in row: c.font=F; c.alignment=Alignment(vertical='top',wrap_text=c.column_letter not in ('A',))
  for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
  ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
H=['Ordem (fácil→difícil)','Nome','Tipo','Categoria no Google','Endereço','Telefone','Site','Nota','Avaliações','Sem site (3)','Só rede social/reserva (2)','<30 avaliações (1)','Nota ≥4,5 c/ presença fraca (2)','Nota <4,0 (−10)','Pontos','Ficha assumida pelo dono?','Motivo','Distância da avenida (m)','Link do Maps']
def line(x,n):
  f=x['f'];r=n+1
  return [x.get('rank',n),x['nome'],x['tipo'],x['cat_google'],x['end'],x['tel'],x['site'],x['nota'],x['aval'],f['sem_site'] or None,f['social'] or None,f['poucas'] or None,f['melhor'] or None,f['cuidado'] or None,
    f'=3*N(J{r})+2*N(K{r})+1*N(L{r})+2*N(M{r})-10*N(N{r})','Conferir no Maps',x['mot'],x['dist'],x['link']]
ws=wb.active; ws.title='Alvos avenida'
sheet(ws,H,[line(x,i) for i,x in enumerate(L,1)],[9,32,16,20,42,16,30,7,10,8,10,9,11,9,8,14,48,10,30])
for i,x in enumerate(L,2):
  fill=BAD if x['f']['cuidado'] else (GOOD if x['f']['melhor'] else None)
  if fill:
    for c in ws[i]: c.fill=fill
  ws.cell(i,19).hyperlink=x['link']
# ordem de visita: segue a avenida do início sul ao fim norte
V=sorted(L,key=lambda x:(x['pos'],x['dist']))
ws2=wb.create_sheet('Ordem de visita')
sheet(ws2,['Parada','Ordem de prioridade','Nome','Tipo','Endereço','Na avenida ou transversal','Metros desde o início da avenida (sul)','Nota','Avaliações','Motivo','Link do Maps'],
  [[i,x['rank'],x['nome'],x['tipo'],x['end'],'Avenida' if x['onav'] else f"Transversal (~{x['dist']} m)",x['pos'],x['nota'],x['aval'],x['mot'],x['link']] for i,x in enumerate(V,1)],[7,10,32,16,42,18,12,7,10,48,30])
for i,x in enumerate(V,2):
  if x['f']['cuidado']:
    for c in ws2[i]: c.fill=BAD
  elif x['f']['melhor']:
    for c in ws2[i]: c.fill=GOOD
  ws2.cell(i,11).hyperlink=x['link']
ws3=wb.create_sheet('Automotivo arredores')
sheet(ws3,H,[line(x,i) for i,x in enumerate(E,1)],[9,32,16,20,42,16,30,7,10,8,10,9,11,9,8,14,48,10,30])
for i,x in enumerate(E,2): ws3.cell(i,1).value=i-1; ws3.cell(i,19).hyperlink=x['link']
ws4=wb.create_sheet('Legenda')
for t in ['Como ler esta planilha',
 'Fonte: Google Maps via SerpApi, buscas feitas em 15/10/2026. Notas, avaliações e telefones mudam: confira na hora da visita.',
 'Área: comércios a até 300 m da Alameda Oscar Niemeyer (a "avenida"), com endereço em Vila da Serra ou Vale do Sereno. Os 300 m pegam o primeiro trecho das ruas transversais (ex.: R. Min. Orozimbo Nonato, Alameda do Ingá, R. do Campo, R. da Paisagem). Ficou de fora quem está na Rodovia (BR-356), que não é transversal.',
 'Pontos (coluna O, fórmula): sem site = 3; site que é só Instagram/Linktree/Facebook ou página de reserva = 2; menos de 30 avaliações (ou nenhuma) = 1; nota 4,5 ou mais com presença fraca = 2 extras (melhor alvo, linha verde); nota abaixo de 4,0 = −10 (vai para o fim, linha laranja).',
 'Ordem: mais pontos primeiro; empate → nota maior; depois mais avaliações (negócio movimentado tende a ter verba).',
 'Nota vazia = o Google não mostra nota (não é zero). Avaliações vazias = nenhuma avaliação.',
 'Ficha assumida pelo dono: a SerpApi não informa esse dado (testado em 15/10). Conferir no Maps: se aparecer "Reivindicar esta empresa", a ficha não é do dono = ótimo gancho de venda.',
 'Ordem de visita: segue a avenida do início sul até o fim norte. Coluna "Ordem de prioridade" mostra quem vale mais a parada.',
 'Aba "Automotivo arredores": estética automotiva e venda de veículos entre 300 m e 1 km, porque na faixa de 300 m quase não há negócios desse tipo.',
 'Dados reais: este arquivo não entra no cofre.']:
  ws4.append([t]); ws4.cell(ws4.max_row,1).font=Font(name='Arial',size=10,bold=(ws4.max_row==1)); ws4.cell(ws4.max_row,1).alignment=Alignment(wrap_text=True)
ws4.column_dimensions['A'].width=120
out='Prospeccao_Av_Oscar_Niemeyer_2026-10-15.xlsx'; wb.save(out)
import json; json.dump({'L':L,'E':E,'V':[x['nome'] for x in V]},open('result.json','w'),ensure_ascii=False,default=str)
print(len(L),len(E))
