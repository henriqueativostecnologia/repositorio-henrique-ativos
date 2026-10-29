import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
SP='dados/'  # pasta com os arquivos de entrada e saída (fica FORA do repositório: dados reais)
UP='dados/'  # planilhas enviadas (fora do repositório)
src=openpyxl.load_workbook(SP+'Visitas_automotivo_e_alimentacao.xlsx')
# segmento e distância dos automotivos, casando pelo link do Maps
seg={};km={}
a=openpyxl.load_workbook(UP+'Alvos_Pampulha_VendaNova_automotivo_2026-10-15.xlsx',read_only=True)
for t in ('Alvos estética','Revendas seminovos'):
    rows=list(a[t].iter_rows(values_only=True)); h={x:i for i,x in enumerate(rows[0])}
    for r in rows[1:]: seg[r[h['Link do Maps']]]=r[h['Segmento']]
m=openpyxl.load_workbook(UP+'Modernizar_site_Pampulha_VendaNova_2026-10-15.xlsx',read_only=True)
rows=list(m['Modernizar site'].iter_rows(values_only=True)); h={x:i for i,x in enumerate(rows[0])}
for r in rows[1:]:
    seg.setdefault(r[h['Link do Maps']],r[h['Segmento']]); km[r[h['Link do Maps']]]=r[h['Distância do Jaraguá (km, linha reta)']]
cols=['Setor','Segmento','Tem site?','Grupo','Prioridade / dia','Nome','Regional','Bairro','Km do Jaraguá (linha reta)','Endereço','Telefone','Site atual','Nota','Avaliações','Situação do site','Argumento para a visita','Link do Maps','Resultado da visita']
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Todos'; ws.append(cols)
out=[]
au=list(src['Automotivo'].iter_rows(values_only=True)); h={x:i for i,x in enumerate(au[0])}
for r in au[1:]:
    if not r[h['Nome']]: continue
    lk=r[h['Link do Maps']]; k=km.get(lk)
    out.append(['Automotivo',seg.get(lk) or 'Automotivo','Não' if r[h['Grupo']].startswith('Sem') else 'Sim',r[h['Grupo']],r[h['Dia / prioridade']],r[h['Nome']],r[h['Regional']],r[h['Bairro']],round(k,1) if k is not None else None,r[h['Endereço']],r[h['Telefone']],r[h['Site atual']],r[h['Nota']],r[h['Avaliações']],r[h['Situação do site']],r[h['Argumento para a visita']],lk,None])
al=list(src['Alimentação'].iter_rows(values_only=True)); h={x:i for i,x in enumerate(al[0])}
for r in al[1:]:
    out.append(['Alimentação',r[h['Segmento']],'Não' if r[h['Grupo']].startswith('Sem') else 'Sim',r[h['Grupo']],r[h['Prioridade']] if r[h['Prioridade']]!='—' else r[h['Faixa de distância']],r[h['Nome']],r[h['Regional']],r[h['Bairro']],r[h['Km do Jaraguá (linha reta)']],r[h['Endereço']],r[h['Telefone']],r[h['Site atual']],r[h['Nota']],r[h['Avaliações']],r[h['Situação do site']],r[h['Argumento para a visita']],r[h['Link do Maps']],None])
out.sort(key=lambda r:((r[6] or 'zz'),(r[7] or '').lower(),r[0],r[5].lower()))
for r in out: ws.append(r)
CAB=PatternFill('solid',fgColor='172631'); VERDE=PatternFill('solid',fgColor='E3F1E6'); AZUL=PatternFill('solid',fgColor='E4EEF7')
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=CAB; c.alignment=Alignment(wrap_text=True,vertical='center')
for row in ws.iter_rows(min_row=2):
    f=VERDE if row[2].value=='Não' else AZUL
    for c in row: c.fill=f; c.alignment=Alignment(wrap_text=True,vertical='top')
    row[12].number_format='0.0'; row[13].number_format='#,##0'; row[8].number_format='0.0'
for i,w in enumerate([12,16,8,22,14,30,12,16,10,38,16,30,6,10,32,46,30,24],1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes='G2'; ws.auto_filter.ref=ws.dimensions
wl=wb.create_sheet('Como ler')
for t in ['Uma aba só com todas as lojas da lista de visitas: automotivo e alimentação, com e sem site.',
 'Para filtrar: clique na setinha do cabeçalho. Exemplos: Bairro = Jaraguá; Setor = Alimentação; Tem site? = Não.',
 'Ordem: por regional e depois por bairro, para as lojas vizinhas ficarem juntas.',
 'Cores: verde = sem site ou só rede social; azul = tem site (modernizar).',
 '"Prioridade / dia": automotivo = dia sugerido ou prioridade do site; alimentação sem site = faixa de distância; alimentação com site = prioridade (Alta, Média, Baixa).',
 'Km do Jaraguá: em branco quando a busca não trouxe a distância (a maioria dos automotivos). Vazio não é zero.',
 'Fontes: Google Maps via SerpApi (15/10/2026) e leitura dos sites pelo Firecrawl (15 e 16/10/2026).']:
    wl.append([t])
wl.column_dimensions['A'].width=130
wb.save(SP+'Visitas_todos_uma_aba.xlsx')
import collections
print(len(out), collections.Counter((r[0],r[2]) for r in out), sum(1 for r in out if r[8] is None), collections.Counter(r[1] for r in out))
