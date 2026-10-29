import openpyxl, copy, sys
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
SP='dados/'  # pasta com os arquivos de entrada e saída (fica FORA do repositório: dados reais)
UP='dados/Alimentacao_Jaragua_Pampulha_2026-10-15.xlsx'  # planilha enviada (fora do repositório)
src=openpyxl.load_workbook(UP)
sem=list(src['Sem site - só rede social'].iter_rows(values_only=True))
com=list(src['Com site (modernizar)'].iter_rows(values_only=True))
H=sem[0]; I={h:i for i,h in enumerate(H)}
wb=openpyxl.load_workbook(SP+'Visitas_automotivo_com-e-sem-site.xlsx')
wb['Lista de visitas'].title='Automotivo'
wb['Como ler'].title='Como ler (automotivo)'
VERDE=PatternFill('solid',fgColor='E3F1E6'); AZUL=PatternFill('solid',fgColor='E4EEF7'); CAB=PatternFill('solid',fgColor='172631')
def faixa(d): return 'Até 2 km' if d<2 else ('2 a 4 km' if d<4 else '4 a 7,5 km')
def site_sit(s):
    if not s: return 'Sem site'
    s=s.lower()
    for k,v in [('instagram','Só Instagram'),('facebook','Só Facebook'),('ifood','Só iFood'),('linktr','Só Linktree'),('beacons','Só página de links (Beacons)'),('wa.me','Só WhatsApp'),('whatsapp','Só WhatsApp')]:
        if k in s: return v
    return 'Só rede social / página de terceiros'
cols=['Grupo','Faixa de distância','Km do Jaraguá (linha reta)','Segmento','Nome','Regional','Bairro','Endereço','Telefone','Site atual','Nota','Avaliações','Situação do site','Prioridade','Argumento para a visita','Link do Maps','Resultado da visita']
ws=wb.create_sheet('Alimentação',1); ws.append(cols)
best=[r for r in sem[1:] if r[I['Melhor alvo']]=='SIM' and r[I['Ativo (10+ aval.)']]=='SIM' and r[I['Regional']] not in ('Barreiro','Centro-Sul')]
best.sort(key=lambda r:r[11])
for r in best:
    obs=r[I['Observação']] or ''
    arg='Página nova (landing page com cardápio, WhatsApp e mapa) + ajuste da ficha do Google'
    if 'rede local' in obs: arg+='. É rede (várias unidades): pedir para falar com o dono da rede'
    if r[I['Regional']]=='A conferir': arg+='. Regional a conferir no mapa'
    ws.append(['Sem site / só rede social',faixa(r[11]),round(r[11],1),r[2],r[3],r[5],r[4],r[6],r[7],r[8] or '—',r[9],r[10],site_sit(r[8]),'—',arg,r[15],None])
# com site: verificados pelo Firecrawl em 15-16/10/2026
C={c[3]:c for c in com[1:]}
ver=[  # EXEMPLOS: (nome igual ao da planilha, prioridade, situação do site, argumento)
 ('Padaria Exemplo - Unidade 1','Alta','Site genérico, sem WhatsApp e sem fotos reais','O site que aparece no Google não fala da padaria de verdade. Confirmar com o dono se foi ele quem fez. É rede: falar com o dono'),
 ('Lanchonete Exemplo','Alta','O link do Google abre o delivery de outra empresa','Cliente que clica em "site" cai em outra empresa. Corrigir a ficha + página própria'),
 ('Confeitaria Exemplo','Alta','Site fora do ar (o endereço não responde)','Quem clica em "site" não chega a lugar nenhum'),
 ('Café Exemplo','Média','Só um Linktree (página de links): sem endereço, horário, fotos ou avaliações','Na prática não tem site: página própria com cardápio'),
 ('Empório Exemplo','Baixa','Site moderno com WhatsApp e fotos; falta endereço, horário, cardápio e avaliações','Ajuste fino: mostrar nota, horário e cardápio'),
]
for nome,pri,sit,arg in ver:
    r=C[nome]
    ws.append(['Com site (modernizar)',faixa(r[11]),round(r[11],1),r[2],r[3],r[5],r[4],r[6],r[7],r[8],r[9],r[10],sit,pri,arg,r[15],None])
# formato
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=CAB; c.alignment=Alignment(wrap_text=True,vertical='center')
for row in ws.iter_rows(min_row=2):
    f=VERDE if row[0].value.startswith('Sem') else AZUL
    for c in row: c.fill=f; c.alignment=Alignment(wrap_text=True,vertical='top')
W=[22,12,10,11,30,12,16,38,16,30,6,10,34,10,46,30,24]
for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes='F2'; ws.auto_filter.ref=ws.dimensions
for i in (11,12):
    for row in ws.iter_rows(min_row=2,min_col=i,max_col=i):
        row[0].number_format='0.0' if i==11 else '#,##0'
# lista completa
wf=wb.create_sheet('Alimentação - lista completa')
for t,rows in (('sem',sem),('com',com)):
    for j,r in enumerate(rows):
        if t=='com' and j==0: continue
        wf.append(list(r))
for c in wf[1]: c.font=Font(bold=True)
wf.freeze_panes='E2'; wf.auto_filter.ref=wf.dimensions
# como ler
wl=wb.create_sheet('Como ler (alimentação)')
for t in [
 f'Alimentação: {sum(1 for r in ws.iter_rows(min_row=2) if r[0].value.startswith("Sem"))} padarias, cafeterias e armazéns sem site ou só com rede social (melhor alvo: nota 4,5 ou mais, 10 ou mais avaliações) + {len(ver)} com site para modernizar.',
 'Ordem: da mais perto para a mais longe do Jaraguá (linha reta). Fora da lista de visitas: Barreiro e Centro-Sul (longe) e as redes grandes (Rede Exemplo A, Rede Exemplo B).',
 'Fontes: Google Maps via SerpApi (15/10/2026) e leitura dos sites pelo Firecrawl (15 e 16/10/2026).',
 'Prioridade dos sites: Alta = fora do ar, link errado ou site que não é da loja; Média = falta o básico ou não abriu na leitura automática; Baixa = site bom, só ajuste fino.',
 '"Site genérico em inglês": páginas .shop que parecem geradas automaticamente. Não afirme ao dono que é golpe: pergunte se foi ele quem fez.',
 'Os demais com site (ainda não abertos) estão na aba "Alimentação - lista completa".',
 'Padaria Exemplo tem várias unidades: é rede. Melhor falar com o dono da rede do que com o gerente da loja.']:
    wl.append([t]); wl.column_dimensions['A'].width=140
out=SP+'Visitas_automotivo_e_alimentacao.xlsx'; wb.save(out); print(out, ws.max_row-1)
