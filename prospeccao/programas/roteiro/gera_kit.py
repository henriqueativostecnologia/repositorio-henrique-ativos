import json, re, html, urllib.parse, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
SP='dados/'  # pasta com os arquivos de entrada e saída (fica FORA do repositório: dados reais)
S={s['Nome']:s for s in json.load(open(SP+'xl/sel.json'))}
def g(p): return next(v for k,v in S.items() if k.startswith(p))
PP='A Padaria Exemplo tem 2 unidades na lista: peça para falar com o dono da rede; é uma conversa só para as 2.'
# EXEMPLOS: (bloco, [(começo do nome igual ao da planilha, argumento de uma linha)])
plano=[
('1. Bairro A (perto de casa)',[
 ('Lava Jato Exemplo','Nota alta e só Instagram. Página com WhatsApp e antes e depois.'),
 ('Padaria Exemplo - Unidade 1','Nota alta, muitas avaliações e sem site. '+PP),
]),
('2. Bairro B',[
 ('Padaria Exemplo - Unidade 2','Nota alta, muitas avaliações e sem site. '+PP),
 ('Café Exemplo','O "site" é só um Linktree: sem endereço, horário, fotos ou avaliações.'),
]),
]
def tel(t): return 'tel:+55'+re.sub(r'\D','',t) if t and t!='None' else None
def maps_dir(addrs): return 'https://www.google.com/maps/dir/'+'/'.join(urllib.parse.quote(a) for a in addrs)
def curto(e): return e.split(' - ')[0]+', '+e.split(' - ')[1].split(',')[0]+', Belo Horizonte' if ' - ' in e else e
linhas=[];n=0;prev='Jaraguá, Belo Horizonte - MG';rotas=[]
for bloco,itens in plano:
    pts=[prev]
    for p,arg in itens:
        s=g(p);n+=1;linhas.append((n,bloco,s,arg));pts.append(curto(s['Endereço']))
    rotas.append((bloco,maps_dir(pts)));prev=pts[-1]
# xlsx
wb=openpyxl.Workbook();ws=wb.active;ws.title='Roteiro 16-10'
cols=['Parada','Bloco','Setor','Nome','Bairro','Endereço','Telefone','Nota','Avaliações','Tem site?','Argumento (uma linha)','Link do Maps','Visitou? (S/N)','Falou com quem','Reação do dono','Objeção','Próximo passo']
ws.append(cols)
for n,b,s,a in linhas: ws.append([n,b,s['Setor'],s['Nome'],s['Bairro'],s['Endereço'],s['Telefone'] if s['Telefone']!='None' else '','%s'%s['Nota'],s['Avaliações'],s['Tem site?'],a,s['Link do Maps'],None,None,None,None,None])
for row in ws.iter_rows(min_row=2):
    row[7].value=float(row[7].value)
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='172631');c.alignment=Alignment(wrap_text=True)
for i,w in enumerate([7,22,11,28,14,34,16,6,10,8,50,28,10,16,30,30,24],1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes='E2'
wr=wb.create_sheet('Rotas no Maps');wr.append(['Bloco','Link da rota (abre no Google Maps com as paradas em ordem)'])
for b,u in rotas: wr.append([b,u])
wr.column_dimensions['A'].width=45;wr.column_dimensions['B'].width=120
wb.save(SP+'Roteiro_visitas_16-10.xlsx')
# html celular
E=html.escape
cards=[];cur=None
for n,b,s,a in linhas:
    if b!=cur:
        if cur: cards.append('</section>')
        u=dict(rotas)[b];cards.append(f'<section><h2>{E(b)}</h2><a class="rota" href="{E(u)}">Abrir a rota deste bloco no Maps</a>');cur=b
    t=tel(s['Telefone'])
    site='Sem site' if s['Tem site?']=='Não' else 'Tem site'
    cards.append(f'''<article><div class="top"><span class="n">{n}</span><div><h3>{E(s["Nome"])}</h3><p class="meta">{E(s["Setor"])} · {E(s["Bairro"])} · nota {format(float(s["Nota"]),".1f").replace(".",",")} ({s["Avaliações"]:,} aval.).replace(",",".") · {site}</p></div></div>
<p class="arg">{E(a)}</p><div class="btns"><a href="{E(s["Link do Maps"])}">Ver no Maps</a>{f'<a href="{t}">Ligar</a>' if t else ''}</div></article>''')
cards.append('</section>')
body='\n'.join(cards).replace(' aval.).replace(",",".")',' aval.)')
# corrige separador de milhar
body=re.sub(r'\((\d{1,3}),(\d{3}) aval\.\)',r'(\1.\2 aval.)',body)
page=f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Visitas 16/10</title>
<style>:root{{--a:#172631;--p:#F5F2EC;--c:#2A5A7E;--l:#DDD7CC;--s:#4B5A63}}*{{box-sizing:border-box}}body{{margin:0;font-family:system-ui,sans-serif;background:var(--p);color:var(--a);line-height:1.45}}
header{{background:var(--a);color:var(--p);padding:18px 16px}}h1{{margin:0;font-size:1.4rem}}header p{{margin:4px 0 0;color:#AFC0CB;font-size:.9rem}}
main{{max-width:640px;margin:0 auto;padding:8px 16px 40px}}h2{{font-size:1.05rem;margin:24px 0 8px}}.rota{{display:block;text-align:center;background:var(--c);color:#fff;padding:12px;border-radius:10px;text-decoration:none;font-weight:700;margin-bottom:10px}}
article{{background:#fff;border:1px solid var(--l);border-radius:12px;padding:12px 14px;margin-bottom:10px}}.top{{display:flex;gap:10px}}.n{{flex:none;width:28px;height:28px;border-radius:50%;background:var(--a);color:#fff;display:grid;place-items:center;font-weight:700;font-size:.85rem}}
h3{{margin:0;font-size:1rem}}.meta{{margin:2px 0 0;color:var(--s);font-size:.82rem}}.arg{{margin:8px 0}}.btns{{display:flex;gap:8px}}.btns a{{flex:1;text-align:center;border:1px solid var(--c);color:var(--c);padding:9px;border-radius:8px;text-decoration:none;font-weight:600}}
.dica{{background:#fff;border-left:4px solid var(--c);border-radius:8px;padding:10px 12px;margin-top:12px;font-size:.92rem}}</style></head><body>
<header><h1>Visitas de 16/10</h1><p>{n} lojas · {len(plano)} blocos · saída do Jaraguá</p></header><main>
<div class="dica"><b>Dica de horário:</b> padaria e café ficam cheios no café da manhã e no almoço. Prefira o meio da manhã e o meio da tarde para eles; lava jato e estética, a qualquer hora.</div>
{body}
<p class="meta" style="margin-top:20px">Fonte: Google Maps (15/10/2026) e leitura dos sites (15 e 16/10/2026). A ordem dentro de cada bloco é por proximidade de bairro; ajuste no Maps se preferir.</p></main></body></html>'''
open(SP+'Roteiro_visitas_16-10.html','w').write(page)
for b,u in rotas: print(b,len(u))
